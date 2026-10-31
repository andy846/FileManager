"""Localization guardrails for Chinese and English UI/output."""

import csv
import json
import os
import re

import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QMessageBox

from app.core.i18n import (
    file_filter, install_qt_translations, remove_qt_translations,
    set_language, tr,
)


class _FakeQ:
    def __init__(self, path):
        self._path = str(path)

    def fileName(self):
        return self._path

    def sync(self):
        pass


class _Settings:
    def __init__(self, path):
        self.data = {}
        self.recent_data = {}
        self.values = {}
        self.q = _FakeQ(path)

    def option(self, name, default):
        return self.data.get(name, default)

    def set_option(self, name, value):
        self.data[name] = value

    def set_options(self, values):
        self.data.update(values)

    def value(self, name, default=None):
        return self.values.get(name, default)

    def set_value(self, name, value):
        self.values[name] = value

    def recent(self, name):
        return list(self.recent_data.get(name, []))

    def remember_path(self, name, path):
        paths = self.recent_data.setdefault(name, [])
        if path in paths:
            paths.remove(path)
        paths.insert(0, path)

    def save_geometry(self, _data):
        pass

    def load_geometry(self):
        return None


def _read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def test_qt_standard_buttons_are_chinese(app):
    set_language("zh_TW")
    translator = install_qt_translations(app)
    try:
        box = QMessageBox()
        box.setStandardButtons(
            QMessageBox.StandardButton.Yes
            | QMessageBox.StandardButton.No
        )
        labels = [button.text() for button in box.buttons()]
        assert any("是" in text for text in labels)
        assert any("否" in text for text in labels)
        assert not any("Yes" in text for text in labels)

        set_language("en")
        install_qt_translations(app)
        english = QMessageBox()
        english.setStandardButtons(
            QMessageBox.StandardButton.Yes
            | QMessageBox.StandardButton.No
        )
        assert all(
            "Yes" in button.text() or "No" in button.text()
            for button in english.buttons()
        )
    finally:
        remove_qt_translations(app)
        set_language("en")
        del translator


def test_file_filters_follow_language():
    set_language("zh_TW")
    try:
        assert file_filter("csv") == "CSV 檔案 (*.csv)"
        assert file_filter("excel") == "Excel 活頁簿 (*.xlsx)"
        assert file_filter("zip") == "ZIP 壓縮檔 (*.zip)"
    finally:
        set_language("en")
    assert file_filter("csv") == "CSV files (*.csv)"
    assert file_filter("excel") == "Excel workbook (*.xlsx)"
    assert file_filter("zip") == "ZIP archive (*.zip)"


def test_cancel_message_is_localized():
    from app.core.workers import CancelledError, OperationControl

    control = OperationControl()
    control.cancel()
    set_language("zh_TW")
    try:
        with pytest.raises(CancelledError) as info:
            control.checkpoint()
        assert str(info.value) == "操作已取消"
        set_language("en")
        assert tr(str(info.value)) == "Operation cancelled"
    finally:
        set_language("en")


def test_remover_deletion_logs_are_fully_localized(tmp_path):
    from openpyxl import load_workbook

    from app.modules.remover.logmanager import LogManager

    settings = _Settings(tmp_path / "settings.ini")
    manager = LogManager(settings)
    try:
        set_language("zh_TW")
        result = manager.save_logs(
            r"C:\test", ["pdf"], [r"C:\test\a.pdf"],
            [(r"C:\test\b.pdf", "locked")],
            "2026-10-26 10:00:00", "2026-10-26 10:01:00",
            size_freed=123,
        )
        rows = _read_csv(result["csv"])
        assert rows[0] == ["狀態", "檔案路徑", "錯誤"]
        assert rows[1][0] == "已刪除"
        assert rows[2][0] == "失敗"
        text = open(result["txt"], encoding="utf-8").read()
        assert "位元組" in text
        assert "bytes" not in text
        workbook = load_workbook(result["xlsx"], read_only=True)
        try:
            assert workbook.sheetnames == ["刪除明細", "摘要"]
            sheet = workbook["刪除明細"]
            assert next(sheet.iter_rows(values_only=True)) == (
                "狀態", "檔案路徑", "錯誤",
            )
        finally:
            workbook.close()

        set_language("en")
        result = manager.save_logs(
            r"C:\test", ["pdf"], [r"C:\test\a.pdf"], [],
            "2026-10-26 10:00:00", "2026-10-26 10:01:00",
            size_freed=123,
        )
        assert _read_csv(result["csv"])[0] == [
            "Status", "File path", "Error",
        ]
        workbook = load_workbook(result["xlsx"], read_only=True)
        try:
            assert workbook.sheetnames == ["Deletion details", "Summary"]
        finally:
            workbook.close()
    finally:
        set_language("en")


def test_remover_scan_exports_use_active_language(
    app, tmp_path, monkeypatch
):
    from openpyxl import load_workbook

    from app.core.models import FileInfo
    from app.modules.remover import page as remover_page

    settings = _Settings(tmp_path / "settings.ini")
    page = remover_page.RemoverPage(settings, None)
    try:
        set_language("zh_TW")
        info = FileInfo(path=str(tmp_path / "a.pdf"), size=128)
        page._reset_counts()
        page._add_row(info)
        page.files = [info]
        page._sel_rows = {0}
        monkeypatch.setattr(
            remover_page.QMessageBox, "information",
            lambda *args, **kwargs: None,
        )
        output_csv = tmp_path / "scan.csv"
        output_xlsx = tmp_path / "scan.xlsx"
        monkeypatch.setattr(
            remover_page.QFileDialog, "getSaveFileName",
            lambda *args, **kwargs: (str(output_csv), ""),
        )
        page._export("csv")
        assert _read_csv(output_csv)[0] == [
            "檔名", "大小", "修改時間", "路徑", "是否選取",
        ]

        monkeypatch.setattr(
            remover_page.QFileDialog, "getSaveFileName",
            lambda *args, **kwargs: (str(output_xlsx), ""),
        )
        page._export("xlsx")
        workbook = load_workbook(output_xlsx, read_only=True)
        try:
            assert workbook.sheetnames == ["掃描結果"]
            sheet = workbook["掃描結果"]
            assert next(sheet.iter_rows(values_only=True)) == (
                "檔名", "大小", "修改時間", "路徑",
            )
        finally:
            workbook.close()
    finally:
        page.deleteLater()
        set_language("en")


def test_settings_language_texts_are_single_language(
    app, tmp_path, monkeypatch
):
    from app.modules.settings import page as settings_page

    settings = _Settings(tmp_path / "settings.ini")
    set_language("zh_TW")
    page = settings_page.SettingsPage(settings, None)
    messages = []
    monkeypatch.setattr(
        settings_page.QMessageBox, "information",
        lambda *args, **kwargs: messages.append(args[2]),
    )
    try:
        assert page.language_combo.itemText(0) == "英文"
        assert page.language_combo.itemText(1) == "繁體中文"
        labels = [
            child.text() for child in page.findChildren(settings_page.QLabel)
        ]
        assert "語言" in labels
        assert not any("Language /" in text for text in labels)

        index = page.language_combo.findData("zh_TW")
        page.language_combo.setCurrentIndex(index)
        page._save()
        restart = [message for message in messages if "重新啟動" in message]
        assert restart
        assert all("Please restart" not in message for message in restart)

        set_language("en")
        assert tr("語言") == "Language"
        assert tr("語言已變更") == "Language changed"
        assert "Please restart" in tr(
            "請重新啟動 FileManager 以套用新語言。"
        )
    finally:
        page.deleteLater()
        set_language("en")


def test_copier_progress_uses_localized_eta(app, tmp_path):
    from app.modules.copier.page import CopierPage

    settings = _Settings(tmp_path / "settings.ini")
    page = CopierPage(settings, None)
    try:
        set_language("zh_TW")
        page._on_stats(1024, 12, 2048)
        assert "剩餘 12 秒" in page.speed_label.text()
        assert "ETA" not in page.speed_label.text()

        set_language("en")
        page._on_stats(1024, 12, 2048)
        assert "ETA 12s" in page.speed_label.text()
    finally:
        page.deleteLater()
        set_language("en")


def test_english_mode_has_no_chinese_ui_chrome(app):
    from app.core.i18n import Localizer
    from app.modules.browser.pane import BrowserProxy
    from app.ui.components import StatusBadge
    from app.ui.main_window import MainWindow

    set_language("en")
    install_qt_translations(app)
    window = MainWindow(_Settings(os.path.join(
        os.path.dirname(__file__), "unused.ini"
    )))
    localizer = Localizer(app, window)
    try:
        window.show()
        localizer.translate_all()
        app.processEvents()
        localizer.translate_all()
        proxy = BrowserProxy()
        assert proxy.headerData(
            0, Qt.Orientation.Horizontal
        ) == "Name"
        assert StatusBadge("running").text() == "Running"
        assert window.nav_items["dashboard"].text(0) == "Dashboard"
    finally:
        localizer.timer.stop()
        app.removeEventFilter(localizer)
        window.close()
        remove_qt_translations(app)
        set_language("en")


def test_english_dynamic_status_strings_are_translated():
    samples = [
        "已開啟壓縮檔：3 個項目",
        "已選取 3 個項目",
        "已剪下 3 個項目",
        "已移至資源回收筒：3 個項目",
        "比對完成：3 個差異",
        "已刪除 3 個項目，失敗 1",
        "分析完成：100 個檔案，10.5 MB",
        "掃描完成: 3 個檔案",
        "清理完成: 3 個檔案",
        "改名完成: 成功 3, 失敗 1",
        "完成: 刪除 3 個, 失敗 1 個",
        "開始複製 3 個檔案",
        "開始搬移 3 個檔案",
        "已加入網路位置：\\\\server\\share",
        "已移除網路位置：\\\\server\\share",
        "正在讀取 \\\\server 的共享...",
        "正在連線 \\\\server\\share ...",
        "同步中 1/3：a.txt",
        "已取消：完成 1 · 跳過 0 · 失敗 1",
        "完成 1 · 跳過 0 · 失敗 1 · 1.0 MB · 1.2s",
        "項目 10 · 差異 3 · 相同 5 · 略過 2 個無法讀取項目",
        "找到 3 個檔案 · 1.2s",
        "成功刪除: 3 個",
        "失敗: 1 個",
        "前 5 個失敗:",
    ]
    set_language("en")
    try:
        cjk = re.compile(r"[\u4e00-\u9fff]")
        for source in samples:
            translated = tr(source)
            assert not cjk.search(translated), (source, translated)
    finally:
        set_language("en")


def test_chinese_placeholder_formatting():
    set_language("zh_TW")
    try:
        assert tr("{count} 個", count="17,094") == "17,094 個"
        assert tr("剩餘 {seconds} 秒", seconds=12) == "剩餘 12 秒"
    finally:
        set_language("en")


def test_english_catalog_has_no_chinese_values():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    path = os.path.join(root, "assets", "i18n", "en.json")
    with open(path, encoding="utf-8") as handle:
        data = json.load(handle)
    cjk = re.compile(r"[\u4e00-\u9fff]")
    assert not [key for key, value in data.items() if cjk.search(value)]
