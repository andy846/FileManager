"""Audit log generation (TXT / CSV / XLSX) for delete & trash operations."""

import os
import csv
import datetime

from app.core.i18n import tr


def _app_data_dir(settings):

    ini = settings.q.fileName()
    base = os.path.dirname(ini)
    return base


class LogManager:
    def __init__(self, settings, log_dir_name="deletion_logs"):
        self.settings = settings
        self.log_dir = os.path.join(_app_data_dir(settings), log_dir_name)
        os.makedirs(self.log_dir, exist_ok=True)

    def generate_timestamp(self):
        return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    def _paths(self, base):
        return (
            os.path.join(self.log_dir, f"{base}.txt"),
            os.path.join(self.log_dir, f"{base}.csv"),
            os.path.join(self.log_dir, f"{base}.xlsx"),
        )

    def save_logs(self, folder, extensions, deleted, failed, started, finished,
                  size_freed=None):
        ts = self.generate_timestamp()
        txt, csvf, xlsxf = self._paths(f"deletion_log_{ts}")
        if size_freed is None:
            # Callers normally pass the exact pre-delete total; files are
            # already gone by now so there is nothing reliable to measure.
            size_freed = 0

        # TXT
        try:
            with open(txt, "w", encoding="utf-8") as f:
                f.write(f"==== {tr('刪除 / 資源回收筒 作業紀錄')} ====\n")
                f.write(f"{tr('開始時間')}: {started}\n")
                f.write(f"{tr('結束時間')}: {finished}\n")
                f.write(f"{tr('目標資料夾')}: {folder}\n")
                f.write(
                    f"{tr('副檔名')}: "
                    f"{', '.join(extensions) or tr('全部')}\n"
                )
                f.write("-" * 40 + "\n")
                for p in deleted:
                    f.write(f"{tr('已刪除')}: {p}\n")
                for p, err in failed:
                    f.write(f"{tr('失敗')}: {p} ({err})\n")
                f.write("-" * 40 + "\n")
                f.write(
                    f"{tr('成功')}: {len(deleted)} | "
                    f"{tr('失敗')}: {len(failed)} | "
                    f"{tr('釋放空間')}: {size_freed} {tr('位元組')}\n"
                )
        except Exception:
            pass

        # CSV
        try:
            with open(csvf, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow([tr("狀態"), tr("檔案路徑"), tr("錯誤")])
                for p in deleted:
                    w.writerow([tr("已刪除"), p, ""])
                for p, err in failed:
                    w.writerow([tr("失敗"), p, err])
        except Exception:
            pass

        # XLSX
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = tr("刪除明細")
            header_fill = PatternFill("solid", fgColor="366092")
            header_font = Font(bold=True, color="FFFFFF")
            green = PatternFill("solid", fgColor="C6EFCE")
            red = PatternFill("solid", fgColor="FFC7CE")
            headers = [tr("狀態"), tr("檔案路徑"), tr("錯誤")]
            for c, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=c, value=h)
                cell.fill = header_fill
                cell.font = header_font
            r = 2
            for p in deleted:
                ws.cell(row=r, column=1, value=tr("已刪除")).fill = green
                ws.cell(row=r, column=2, value=p)
                r += 1
            for p, err in failed:
                ws.cell(row=r, column=1, value=tr("失敗")).fill = red
                ws.cell(row=r, column=2, value=p)
                ws.cell(row=r, column=3, value=err)
                r += 1
            for col in ("A", "B", "C"):
                ws.column_dimensions[col].width = 60 if col != "A" else 12

            summary = wb.create_sheet(tr("摘要"))
            rows = [
                (tr("開始時間"), started),
                (tr("結束時間"), finished),
                (tr("目標資料夾"), folder),
                (tr("副檔名"), ", ".join(extensions) or tr("全部")),
                (tr("成功數量"), len(deleted)),
                (tr("失敗數量"), len(failed)),
                (tr("釋放空間 (位元組)"), size_freed),
            ]
            for i, (k, v) in enumerate(rows, 1):
                summary.cell(row=i, column=1, value=k).font = Font(bold=True)
                summary.cell(row=i, column=2, value=v)
            summary.column_dimensions["A"].width = 22
            summary.column_dimensions["B"].width = 60
            wb.save(xlsxf)
        except Exception:
            xlsxf = None

        return {"txt": txt, "csv": csvf, "xlsx": xlsxf,
                "size_freed": size_freed}
